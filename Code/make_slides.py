# 生成 skill demonstration 的 PPT（极简：白底黑字，Arial）
# 用法：python make_slides.py   （需要 python-pptx）
import re
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BASE = Path(__file__).resolve().parent.parent      # skill demonstration 文件夹
TEMPLATE = Path.home() / 'Downloads' / 'Skill_Demonstration_Template.pptx'
OUT = BASE / 'Skill_Demo_Team3.pptx'

BLACK, GRAY, LIGHT = RGBColor(0, 0, 0), RGBColor(0x59, 0x59, 0x59), RGBColor(0xBF, 0xBF, 0xBF)
FONT = 'Arial'

prs = Presentation(TEMPLATE)                        # 用课程模板的页面尺寸（13.33 x 7.5 英寸）
# 删掉模板里原有的 5 页，只保留版式
for sid in list(prs.slides._sldIdLst):
    prs.part.drop_rel(sid.rId); prs.slides._sldIdLst.remove(sid)
BLANK = [l for l in prs.slide_layouts if l.name == 'Blank'][0]


def text(slide, x, y, w, h, runs, size=20, color=BLACK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=0):
    # runs: 一个字符串 = 一段；列表 = 多段
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate([runs] if isinstance(runs, str) else runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        r = p.add_run(); r.text = line
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    return box


def title(slide, s):
    text(slide, 0.6, 0.45, 12.1, 0.8, s, size=36, bold=True)


def circle(slide, cx, cy, r, label='', size=14, line=BLACK):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
    c.fill.background(); c.line.color.rgb = line; c.line.width = Pt(1.25); c.shadow.inherit = False
    if label:
        tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        r_ = tf.paragraphs[0].add_run(); r_.text = label
        r_.font.name, r_.font.size, r_.font.color.rgb = FONT, Pt(size), BLACK
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return c


def line(slide, x1, y1, x2, y2, color=BLACK, width=1.25, arrow=False):
    l = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    l.line.color.rgb = color; l.line.width = Pt(width)
    if arrow:                                       # python-pptx 没有箭头接口，直接写 XML
        ln = l.line._get_or_add_ln()
        ln.append(ln.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd', {'type': 'triangle'}))
    return l


# ---------- 公式：用文字拼出来（字母斜体，_2 = 下标） ----------
MATH = 'Times New Roman'


def math(slide, x, y, w, h, s, size=24, align=PP_ALIGN.CENTER):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align
    for tok in re.findall(r'[_^]\w|[A-Za-z]+|.', s, re.S):
        r = p.add_run(); r.text = tok[1:] if tok[0] in '_^' else tok
        r.font.name, r.font.size, r.font.color.rgb = MATH, Pt(size), BLACK
        r.font.italic = tok.isalpha() and tok != 'T'    # 变量用斜体，∂、数字、转置 T 用正体
        if tok[0] in '_^':
            r.font._element.set('baseline', '-25000' if tok[0] == '_' else '30000')   # _ 下标，^ 上标
        if tok in '⊙Σ':
            r.font.name = 'Cambria Math'                # Times New Roman 没有 ⊙
    return box


def frac(slide, cx, cy, top, bottom, w=0.95, size=24):
    # 分数：上面 top，中间一条横线，下面 bottom；(cx, cy) 是横线中心
    math(slide, cx - w / 2, cy - 0.5, w, 0.45, top, size)
    line(slide, cx - w / 2 + 0.05, cy, cx + w / 2 - 0.05, cy, width=1)
    math(slide, cx - w / 2, cy + 0.05, w, 0.45, bottom, size)


def dot(slide, cx, cy, r=0.06):
    d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
    d.fill.solid(); d.fill.fore_color.rgb = BLACK; d.line.fill.background(); d.shadow.inherit = False


# ---------- 第 1 页：标题 ----------
s = prs.slides.add_slide(BLANK)
text(s, 0.9, 1.5, 11.5, 1.6, ['Backpropagation vs.', 'Feedback Alignment'], size=44, bold=True)
text(s, 0.9, 3.35, 11.5, 0.6, 'Teaching a network we built ourselves to read handwritten digits', size=22, color=GRAY)
text(s, 0.9, 4.9, 11.5, 1.0, ['Team 3', 'Wenrui Chen, Chengweiran Liu, Xianchen Fang'], size=20, space=4)
text(s, 0.9, 6.6, 11.5, 0.4, 'EN.553.696 Methods in Computational Neuroscience', size=14, color=GRAY)

# ---------- 第 2 页：团队 ----------
s = prs.slides.add_slide(BLANK)
title(s, 'Team Composition')
team = [('WC', 'Wenrui Chen', 'Biomedical Engineering'),
        ('CL', 'Chengweiran Liu', 'Biomedical Engineering'),
        ('XF', 'Xianchen Fang', 'Applied Mathematics & Statistics')]
for i, (ini, name, major) in enumerate(team):
    cx = 2.2 + i * 4.47
    circle(s, cx, 3.0, 0.75, ini, size=24)
    text(s, cx - 2.0, 4.1, 4.0, 0.5, name, size=24, bold=True, align=PP_ALIGN.CENTER)
    text(s, cx - 2.0, 4.65, 4.0, 0.9, [major, "Master's, Year 1"], size=18, color=GRAY, align=PP_ALIGN.CENTER, space=2)
s.notes_slide.notes_text_frame.text = 'About 1 min.'

# ---------- 第 3 页：Introduction & Motivation ----------
s = prs.slides.add_slide(BLANK)
title(s, 'Skill Demo: Introduction & Motivation')
text(s, 0.6, 1.55, 7.0, 0.5, 'What we do', size=24, bold=True)
text(s, 0.6, 2.1, 7.0, 2.0,
     ['1.  Train a neural network with backpropagation to recognize handwritten digits (MNIST, a public dataset)',
      '2.  Feedback alignment: what it is, and how it compares with backpropagation on the same network'],
     size=20, space=10)
text(s, 0.6, 4.25, 7.0, 0.5, 'Why', size=24, bold=True)
text(s, 0.6, 4.8, 7.0, 2.2,
     ['In class: backpropagation on a two-neuron chain',
      'Can the same idea work on a real problem?',
      'With only what we learned in class, nothing extra, we built and trained our own network'],
     size=20, space=10)

# 右侧示意图：课堂上的链  ->  我们的网络
X0 = 8.4
text(s, X0, 1.6, 4.3, 0.4, 'In class', size=16, color=GRAY)
for j, lab in enumerate(['s0', 's1', 's2']):
    circle(s, X0 + 0.45 + j * 1.6, 2.45, 0.3, lab, size=14)
    if j < 2:
        line(s, X0 + 0.75 + j * 1.6, 2.45, X0 + 1.75 + j * 1.6, 2.45, arrow=True)
        text(s, X0 + 0.75 + j * 1.6, 2.0, 1.0, 0.3, f'w{j}', size=14, align=PP_ALIGN.CENTER)
line(s, X0 + 2.05, 3.05, X0 + 2.05, 3.6, color=GRAY, arrow=True)
text(s, X0, 3.7, 4.3, 0.4, 'Our network', size=16, color=GRAY)
cols = [(X0 + 0.45, 5, '784', 'pixels'), (X0 + 2.05, 6, '1000', 'hidden'), (X0 + 3.65, 4, '10', 'digits')]
pos = []
for x, n, _, _ in cols:                             # 每一层画几个圆代表，中间用 ⋮ 表示省略
    ys = [4.45 + k * 0.32 for k in range(n)]
    pos.append([(x, y) for y in ys])
for a, b in zip(pos[:-1], pos[1:]):                 # 层与层之间全连接
    for (x1, y1) in a:
        for (x2, y2) in b:
            line(s, x1 + 0.12, y1, x2 - 0.12, y2, color=LIGHT, width=0.5)
for (x, n, num, lab), p in zip(cols, pos):
    for (cx, cy) in p:
        circle(s, cx, cy, 0.12)
    text(s, x - 0.7, 6.45, 1.4, 0.6, [num, lab], size=14, align=PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text = (
    'About 1 min.\n'
    'Our demo has two parts. First, we train a neural network with backpropagation to recognize handwritten digits '
    'from MNIST, a public dataset. Second, we explain feedback alignment, apply it to the same network, '
    'and compare it with backpropagation.\n'
    'Why this project? In class we learned backpropagation on a very small model: a chain of two neurons. '
    'We wanted to know how to take that into real life. It turns out that with only what we learned in class, '
    'nothing extra, we can build our own network and train a working model. That was exciting for us: '
    'using what we learned in the classroom on a real problem.')

# ---------- 第 4 页：Quick review（课上的 backprop） ----------
s = prs.slides.add_slide(BLANK)
title(s, 'Quick Review: Backpropagation in Class')
text(s, 0.6, 1.55, 12.1, 0.5, 'What does it train?', size=24, bold=True)
text(s, 0.6, 2.1, 12.1, 0.9,
     'Given a known input s0 and a known output, train the weights w0, w1 so that the output of the network matches the known output',
     size=20)
# 课上的链：s0 —• (x1) — s1 —• (x2) — s2 —• out
X, Y = 3.2, 3.55                                     # 链的起点和高度
math(s, X - 0.5, Y - 0.25, 0.45, 0.5, 's_0', align=PP_ALIGN.RIGHT)
starts, ends = [0, 1.99, 4.54], [1.25, 3.8, 6.35]   # 三段线（相对 X），每段末尾一个点
for k in range(3):
    line(s, X + starts[k], Y, X + ends[k], Y, width=2)
    dot(s, X + ends[k], Y)
    if k < 2:
        math(s, X + ends[k] - 0.55, Y - 0.62, 0.6, 0.45, f'w_{k}')         # 权重写在点的左上方
        circle(s, X + ends[k] + 0.4, Y, 0.34)
        math(s, X + ends[k] + 0.06, Y - 0.25, 0.68, 0.5, f'x_{k+1}')      # 神经元里写 x1 / x2
    if k > 0:
        mid = X + (starts[k] + ends[k]) / 2
        math(s, mid - 0.3, Y - 0.62, 0.6, 0.45, f's_{k}')                 # 神经元输出 s1 / s2
math(s, X + 6.6, Y - 0.25, 0.8, 0.5, 'out', align=PP_ALIGN.LEFT)

text(s, 0.6, 4.35, 12.1, 0.5, 'How? The chain rule', size=24, bold=True)
# 链式法则：∂C/∂w0 = ∂C/∂out · ∂out/∂s2 · ∂s2/∂x2 · ∂x2/∂s1 · ∂s1/∂x1 · ∂x1/∂w0
EY = 5.6                                             # 公式横线的高度
frac(s, 1.1, EY, '∂C', '∂w_0')
math(s, 1.6, EY - 0.25, 0.4, 0.5, '=')
for k, (t, b) in enumerate([('∂C', '∂out'), ('∂out', '∂s_2'), ('∂s_2', '∂x_2'), ('∂x_2', '∂s_1'), ('∂s_1', '∂x_1'), ('∂x_1', '∂w_0')]):
    frac(s, 2.45 + k * 0.95, EY, t, b, w=0.9)
# 箭头 -> 更新权重：w0 ← w0 − α ∂C/∂w0
line(s, 7.95, EY, 9.0, EY, arrow=True)
text(s, 7.95, EY - 0.45, 1.05, 0.35, 'update', size=14, color=GRAY, align=PP_ALIGN.CENTER)
math(s, 8.75, EY - 0.25, 2.3, 0.5, 'w_0 ← w_0 − α', align=PP_ALIGN.RIGHT)
frac(s, 11.55, EY, '∂C', '∂w_0', w=0.9)

text(s, 0.6, 6.6, 12.1, 0.5, '... so how do we use this to recognize an image?', size=22)
s.notes_slide.notes_text_frame.text = (
    'Quick review of what we learned in class.\n'
    'What does backpropagation train? We are given a known input and a known output. '
    'We train the weight of each neuron so that, starting from the input, the network produces an output that matches the known output.\n'
    'How do we train it? With the chain rule: we follow the error back through the chain, one step at a time, '
    'to find how the cost C changes when we change w0. Then we use it to update the weight: '
    'w0 moves a small step, of size alpha, in the direction that makes the cost C smaller.\n'
    'So how do we use this to recognize an image?')

# ---------- 第 5 页：输入和输出 ----------
import gzip
import numpy as np
from PIL import Image
with gzip.open(BASE / 'Data' / 'MNIST' / 'train-images-idx3-ubyte.gz') as f:
    img = np.frombuffer(f.read(), np.uint8, offset=16).reshape(-1, 28, 28)[2]   # 训练集第 3 张，是一个 4
Image.fromarray(255 - img).resize((280, 280), Image.NEAREST).save(BASE / 'Figures' / 'slide_digit4.png')  # 白底黑字
flat = img.reshape(-1) / 255.0


def vec(slide, x, y, cells, w=0.75, h=0.36, labels=None, bold=None):
    # 画一个列向量：cells 里每个字符串是一格，'⋮' 表示省略
    for i, c in enumerate(cells):
        cy = y + i * h
        if c == '⋮':
            text(slide, x, cy, w, h, '⋮', size=16, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        else:
            r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(cy), Inches(w), Inches(h))
            r.fill.background(); r.line.color.rgb = BLACK; r.line.width = Pt(1); r.shadow.inherit = False
            text(slide, x, cy, w, h, c, size=16, bold=(i == bold), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if labels:
            text(slide, x - 0.45, cy, 0.35, h, labels[i], size=14, color=GRAY, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


s = prs.slides.add_slide(BLANK)
title(s, 'Input and Output')
VY = 1.95                                            # 两个列向量的顶部
# 1) 28 x 28 的图
pic = s.shapes.add_picture(str(BASE / 'Figures' / 'slide_digit4.png'), Inches(0.6), Inches(2.25), width=Inches(2.6))
pic.line.color.rgb = LIGHT; pic.line.width = Pt(1)            # 细灰框，看得出 28 x 28 的边界
text(s, 0.6, 5.95, 2.6, 0.8, ['Handwritten digit', '28 × 28 pixels'], size=16, align=PP_ALIGN.CENTER, space=2)
line(s, 3.45, 3.55, 4.35, 3.55, arrow=True)
text(s, 3.45, 3.1, 0.9, 0.35, 'unroll', size=14, color=GRAY, align=PP_ALIGN.CENTER)
# 2) 输入：784 x 1 列向量（真实像素值，第 161~163 个像素）
vals = [f'{v:.1f}' for v in flat[160:163]]
vec(s, 4.55, VY, ['0.0', '0.0', '⋮', *vals, '⋮', '0.0', '0.0'], h=0.4)
text(s, 3.9, 5.95, 2.05, 0.8, ['Input: 784 × 1', 'each value 0 to 1'], size=16, align=PP_ALIGN.CENTER, space=2)
# 3) 网络：784 -> 1000 -> 10
line(s, 5.55, 3.55, 6.35, 3.55, arrow=True)
cols = [(6.85, 5), (8.05, 6), (9.25, 4)]
pos = [[(x, 3.55 + (k - (n - 1) / 2) * 0.34) for k in range(n)] for x, n in cols]
for a, b in zip(pos[:-1], pos[1:]):
    for (x1, y1) in a:
        for (x2, y2) in b:
            line(s, x1 + 0.12, y1, x2 - 0.12, y2, color=LIGHT, width=0.5)
for p in pos:
    for (cx, cy) in p:
        circle(s, cx, cy, 0.12)
for (x, _), num in zip(cols, ['784', '1000', '10']):
    text(s, x - 0.6, 4.75, 1.2, 0.35, num, size=14, align=PP_ALIGN.CENTER)
text(s, 6.35, 5.95, 3.4, 0.8, ['Network', '1 hidden layer of 1000 neurons'], size=16, align=PP_ALIGN.CENTER, space=2)
line(s, 9.75, 3.55, 10.55, 3.55, arrow=True)
# 4) 理想输出：10 x 1，只有 4 那一格是 1
vec(s, 11.25, VY, ['1' if d == 4 else '0' for d in range(10)], h=0.36,
    labels=[str(d) for d in range(10)], bold=4)
text(s, 10.3, 5.95, 2.65, 0.8, ['Ideal output: 10 × 1', '1 at the true digit'], size=16, align=PP_ALIGN.CENTER, space=2)
s.notes_slide.notes_text_frame.text = (
    'Input: each handwritten digit is a 28 by 28 grayscale image. We unroll it row by row into one long column '
    'of 784 numbers, each between 0 and 1. This column is s0, the input.\n'
    'Ideal output: a column of 10 numbers, one for each digit 0 to 9. For an image of a 4, the ideal output is 1 '
    'at position 4 and 0 everywhere else.\n'
    "In between: one hidden layer with 1000 neurons. The network's guess is the output neuron with the largest value.")

# ---------- 第 6 页：很多神经元时怎么更新权重 ----------
s = prs.slides.add_slide(BLANK)
title(s, 'Many Neurons: How to Update a Weight')
def ten_paths(s, show_w1=False):
    # 示意图：s0[1] -> s1[1] -> 全部 10 个输出，一共 10 条路（第 6 页和第 8、9 页共用）
    CX = [1.25, 2.95, 4.65]
    ys_in = [1.95 + k * 0.42 for k in range(5)]
    ys_hid = [1.9 + k * 0.35 for k in range(6)]
    ys_out = [1.8 + k * 0.22 for k in range(10)]
    for i, y1 in enumerate(ys_in):                       # 先画灰色的线
        for j, y2 in enumerate(ys_hid):
            if (i, j) != (0, 0):
                line(s, CX[0] + 0.12, y1, CX[1] - 0.12, y2, color=LIGHT, width=0.5)
    for j, y1 in enumerate(ys_hid[1:]):
        for y2 in ys_out:
            line(s, CX[1] + 0.12, y1, CX[2] - 0.1, y2, color=LIGHT, width=0.5)
    line(s, CX[0] + 0.12, ys_in[0], CX[1] - 0.12, ys_hid[0], width=2.25)    # 再画黑色的 10 条路
    for y2 in ys_out:
        line(s, CX[1] + 0.12, ys_hid[0], CX[2] - 0.1, y2, width=1.5)
    for k, y in enumerate(ys_in):
        circle(s, CX[0], y, 0.12, line=BLACK if k == 0 else LIGHT)
    for k, y in enumerate(ys_hid):
        circle(s, CX[1], y, 0.12, line=BLACK if k == 0 else LIGHT)
    for y in ys_out:
        circle(s, CX[2], y, 0.1)
    math(s, 0.2, ys_in[0] - 0.2, 0.85, 0.4, 's_0[1]', size=18, align=PP_ALIGN.RIGHT)
    math(s, CX[1] - 0.5, ys_hid[0] - 0.6, 1.0, 0.4, 's_1[1]', size=18)
    math(s, CX[0] + 0.3, ys_in[0] - 0.5, 1.1, 0.4, 'W_0[1,1]', size=18)
    math(s, CX[2] + 0.2, ys_out[0] - 0.2, 0.9, 0.4, 's_2[1]', size=18, align=PP_ALIGN.LEFT)
    math(s, CX[2] + 0.2, ys_out[-1] - 0.2, 0.9, 0.4, 's_2[10]', size=18, align=PP_ALIGN.LEFT)
    if show_w1:                                      # 在 s1 -> s2 下面画一个括号，标出 W1 是这一整层的权重
        BY = 4.0
        line(s, CX[1], BY, CX[2], BY)
        line(s, CX[1], BY - 0.08, CX[1], BY); line(s, CX[2], BY - 0.08, CX[2], BY)
        math(s, CX[1], BY + 0.02, CX[2] - CX[1], 0.35, 'W_1', size=18)
        text(s, CX[1] - 0.3, BY + 0.35, CX[2] - CX[1] + 0.6, 0.3, 'all weights from s1 to s2', size=12,
             color=GRAY, align=PP_ALIGN.CENTER)


ten_paths(s)
# 右边文字
text(s, 6.3, 1.75, 6.4, 2.6,
     ['Chain: one path from w0 to C',
      'Our network: W0[1,1] changes s1[1], and s1[1] feeds all 10 output neurons',
      'So there are 10 paths. Write each one with the chain rule and add them up'],
     size=20, space=12)
# 10 条路相加
EY = 4.8
frac(s, 1.15, EY, '∂C', '∂W_0[1,1]', w=1.2, size=22)
math(s, 1.8, EY - 0.25, 0.35, 0.5, '=', size=22)
math(s, 2.15, EY - 0.35, 0.5, 0.7, 'Σ', size=36)
math(s, 2.1, EY + 0.3, 0.6, 0.3, 'k=1', size=12)
math(s, 2.1, EY - 0.6, 0.6, 0.3, '10', size=12)
for k, (t, b) in enumerate([('∂C', '∂s_2[k]'), ('∂s_2[k]', '∂x_2[k]'), ('∂x_2[k]', '∂s_1[1]'),
                            ('∂s_1[1]', '∂x_1[1]'), ('∂x_1[1]', '∂W_0[1,1]')]):
    frac(s, 3.4 + k * 1.3, EY, t, b, w=1.2, size=22)
# 矩阵形式
text(s, 0.6, 5.6, 12.1, 0.45, '... but writing every path is tedious, so we write it with matrices:', size=20)
MY = 6.6
frac(s, 1.15, MY, '∂C', '∂W_0', w=0.9, size=22)
math(s, 1.65, MY - 0.25, 7.4, 0.5,
     '= [ ( W_1^T [ (s_2 − t) ⊙ s_2 ⊙ (1 − s_2) ] ) ⊙ s_1 ⊙ (1 − s_1) ] s_0^T', size=22, align=PP_ALIGN.LEFT)
text(s, 9.6, MY - 0.45, 3.2, 0.9, ['t = ideal output', '⊙ = multiply element by element'], size=14, color=GRAY, space=2)
s.notes_slide.notes_text_frame.text = (
    'Now, how do we update the weights when there are many neurons? In a chain, there is only one path from w0 to C, '
    'so we just follow the chain.\n'
    'In our network, take the weight from the first input neuron to the first hidden neuron, W0[1,1]. '
    'Changing it changes s1[1], and s1[1] feeds all 10 output neurons. So there are 10 paths from W0[1,1] to C. '
    'We write each path with the chain rule, exactly like before, and add the 10 paths together.\n'
    'Writing every path for every weight is tedious, so we write the same thing with matrices. '
    'The circle-dot means multiply element by element, and t is the ideal output.')

# ---------- 第 7 页：结果（backprop） ----------
from pptx.chart.data import XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE
from pptx.oxml.ns import qn
bp_acc = 100 * (1 - np.load(BASE / 'Results' / 'result_bp.npz')['err'])   # 第 0 个是训练前
s = prs.slides.add_slide(BLANK)
title(s, 'Result: Backpropagation')
# 左边：每个 epoch 之后在测试集上的正确率（原生图表，可以在 PowerPoint 里改）
text(s, 0.6, 1.45, 6.2, 0.4, 'Test accuracy after each epoch', size=18, bold=True)
cd = XyChartData(); ser = cd.add_series('backprop')
for ep in range(1, len(bp_acc)):
    ser.add_data_point(ep, round(float(bp_acc[ep]), 2))
ch = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, Inches(0.45), Inches(1.85), Inches(6.3), Inches(3.0), cd).chart
ch.has_legend = False; ch.font.name, ch.font.size, ch.font.color.rgb = FONT, Pt(12), GRAY
pl = ch.plots[0].series[0]
pl.format.line.color.rgb = BLACK; pl.format.line.width = Pt(2); pl.smooth = False
pl.marker.style = XL_MARKER_STYLE.CIRCLE; pl.marker.size = 6
pl.marker.format.fill.solid(); pl.marker.format.fill.fore_color.rgb = BLACK; pl.marker.format.line.color.rgb = BLACK
xa, ya = ch.category_axis, ch.value_axis
xa.minimum_scale, xa.maximum_scale, xa.major_unit = 0, 20, 5
ya.minimum_scale, ya.maximum_scale, ya.major_unit = 92, 100, 2
ya.has_major_gridlines = True; ya.major_gridlines.format.line.color.rgb = RGBColor(0xE6, 0xE6, 0xE6)
xa.has_major_gridlines = False
for ax, lab in [(xa, 'epoch'), (ya, 'accuracy (%)')]:
    ax.format.line.color.rgb = GRAY
    ax.has_title = True; ax.axis_title.text_frame.text = lab
    r = ax.axis_title.text_frame.paragraphs[0].runs[0]
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(12), False, GRAY
# 右边：训练和测试，一句话一个
text(s, 7.3, 1.45, 5.4, 1.0, ['Train', '60,000 images × 20 epochs', '= 1.2 million weight updates'], size=18, space=2)
s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
text(s, 7.3, 2.75, 5.4, 1.0, ['Test', '10,000 images it has never seen'], size=18, space=2)
s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
text(s, 7.3, 3.7, 5.4, 1.0, f'{bp_acc[-1]:.2f}% correct', size=48, bold=True)
# 下面：网络在测试图上的猜测（从 fig4 裁出来：去掉标题行和 FA 那一行，压掉图和文字之间的空白）
g = Image.open(BASE / 'Figures' / 'fig4_example_guesses.png').convert('RGB')
top, bp_row = g.crop((0, 140, g.width, 560)), g.crop((0, 665, g.width, 745))
crop = Image.new('RGB', (g.width, top.height + bp_row.height), 'white')
crop.paste(top, (0, 0)); crop.paste(bp_row, (0, top.height))
crop.save(BASE / 'Figures' / 'slide_guesses_bp.png')
gw = 11.0
s.shapes.add_picture(str(BASE / 'Figures' / 'slide_guesses_bp.png'), Inches((13.333 - gw) / 2), Inches(5.2), width=Inches(gw))
s.notes_slide.notes_text_frame.text = (
    'Now the result. We trained for 20 epochs: the network sees all 60,000 training images 20 times, '
    'and we update the weights after every image. That is 1.2 million updates.\n'
    'Then we test it on 10,000 images it has never seen. It gets 98.33% of them right.\n'
    'Here are some of its guesses. Most are right, like this 7, 2, 1, 0. The two it gets wrong are hard ones: '
    'this 4 looks a lot like a 2, and this 6 is close to a 0.\n'
    'For a network built only from what we learned in class, that is a pretty good number.')

# ---------- 第 8、9 页：大脑能这么做吗？-> 换成随机矩阵 B ----------
# 两页内容一样，第 9 页多了"划掉 W1ᵀ 换成 B"，翻页时就像动画
from pptx.enum.dml import MSO_LINE


def fa_intro(swap):
    s = prs.slides.add_slide(BLANK)
    title(s, "... Let's Think a Little Further")
    ten_paths(s, show_w1=True)
    # 矩阵公式，拆成三段，方便单独划掉中间的 W1ᵀ
    MY = 4.9
    frac(s, 0.95, MY, '∂C', '∂W_0', w=0.8, size=18)
    math(s, 1.15, MY - 0.25, 0.75, 0.5, '= [ (', size=18, align=PP_ALIGN.RIGHT)
    math(s, 1.9, MY - 0.25, 0.55, 0.5, 'W_1^T', size=18)
    math(s, 2.45, MY - 0.25, 5.0, 0.5, '[ (s_2 − t) ⊙ s_2 ⊙ (1 − s_2) ] ) ⊙ s_1 ⊙ (1 − s_1) ] s_0^T', size=18, align=PP_ALIGN.LEFT)
    # 右边：大脑能这么做吗？画两个神经元
    text(s, 7.7, 1.5, 5.0, 0.5, 'Can the brain do this?', size=24, bold=True)
    NY = 2.9                                         # 神经元的高度
    circle(s, 8.55, NY, 0.3)                         # 神经元 s1
    for (x1, y1, x2, y2) in [(8.27, -0.13, 7.85, -0.5), (8.27, 0.13, 7.85, 0.5), (8.25, 0, 7.75, 0),
                             (7.85, -0.5, 7.7, -0.65), (7.85, -0.5, 7.75, -0.35), (7.85, 0.5, 7.7, 0.65), (7.85, 0.5, 7.75, 0.35)]:
        line(s, x1, NY + y1, x2, NY + y2)            # 树突（y 是相对 NY 的）
    line(s, 8.85, NY, 11.0, NY, width=1.5)           # 轴突
    dot(s, 11.05, NY, r=0.07)                        # 轴突末端
    line(s, 11.2, NY, 11.55, NY)                     # 下一个神经元的树突（中间留一个空隙 = 突触）
    circle(s, 11.85, NY, 0.3)                        # 神经元 s2
    math(s, 8.25, NY - 0.75, 0.6, 0.4, 's_1', size=18)                # 名字写在神经元上方
    math(s, 11.55, NY - 0.75, 0.6, 0.4, 's_2', size=18)
    text(s, 10.6, NY + 0.15, 1.05, 0.3, 'synapse', size=12, color=GRAY, align=PP_ALIGN.CENTER)
    line(s, 9.2, NY - 0.45, 10.75, NY - 0.45, arrow=True)          # 正向：信号沿轴突过去
    text(s, 9.2, NY - 0.85, 1.55, 0.35, 'signal', size=14, align=PP_ALIGN.CENTER)
    FB = NY + 0.75                                   # 反向：一条单独的 feedback 连接，从 s2 绕回 s1
    line(s, 11.85, NY + 0.3, 11.85, FB, color=GRAY)
    line(s, 11.85, FB, 8.55, FB, color=GRAY)
    line(s, 8.55, FB, 8.55, NY + 0.32, color=GRAY, arrow=True)
    text(s, 9.2, FB - 0.3, 1.55, 0.28, 'feedback connection', size=12, color=GRAY, align=PP_ALIGN.CENTER)
    t = text(s, 8.55, FB + 0.15, 4.2, 0.7, ['error from s2 can be sent back  ✓', 'its weights equal W1 exactly  ✗'],
             size=16, space=2)
    t.text_frame.paragraphs[1].runs[0].font.bold = True
    text(s, 7.7, 4.6, 5.0, 1.3,
         ['Hard.', 'That feedback connection would have to match W1 exactly, and keep matching it as W1 learns.'],
         size=18, space=4)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    if swap:
        line(s, 1.96, MY + 0.12, 2.4, MY - 0.12, width=2)          # 划掉 W1ᵀ
        math(s, 1.9, MY - 0.75, 0.55, 0.45, 'B', size=22)
        s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
        t = text(s, 0.6, 6.35, 12.1, 0.5, 'Feedback alignment: ', size=22, bold=True)
        r1 = t.text_frame.paragraphs[0].add_run(); r1.text = 'replace W1ᵀ with B, a fixed random matrix that is never trained'
        r1.font.name, r1.font.size, r1.font.color.rgb = FONT, Pt(22), BLACK
        text(s, 0.6, 6.95, 12.1, 0.35,                 # 文献
             'Lillicrap et al. (2016). Random synaptic feedback weights support error backpropagation '
             'for deep learning. Nature Communications, 7, 13276.', size=11, color=GRAY)
    return s


s = fa_intro(False)
s.notes_slide.notes_text_frame.text = (
    "Now let's think a little further. Look at the formula again. To send the error back to the hidden layer, "
    'we use W1 transpose: the same weights the signal used going forward.\n'
    'Can the brain do this? Sending the error back is possible: the brain has many feedback connections, '
    'so a separate connection can carry the error from s2 back to s1.\n'
    'The hard part is the weights. That feedback connection would need exactly the same weights as W1, '
    'and keep them matched every time W1 learns. Two separate connections have no way to know each other\'s weights.')
s = fa_intro(True)
s.notes_slide.notes_text_frame.text = (
    'So what if we just do not use W1 transpose? Feedback alignment replaces it with B, '
    'a fixed random matrix. B is set once at the start and never trained. '
    'In our code, this is a one-line change.')

# ---------- 第 10 页：结果（feedback alignment） ----------
acc = {m: 100 * (1 - np.load(BASE / 'Results' / f'result_{m}.npz')['err']) for m in ('bp', 'fa', 'shallow')}
# 每种方法的线条样式：(名字, 颜色, 虚线)
style = {'bp': ('Backprop', BLACK, False), 'fa': ('Feedback alignment', GRAY, True),
         'shallow': ('Last layer only', LIGHT, False)}
s = prs.slides.add_slide(BLANK)
title(s, 'Result: Feedback Alignment')
text(s, 0.6, 1.45, 6.2, 0.4, 'Test accuracy after each epoch', size=18, bold=True)
cd = XyChartData()
for m in ('shallow', 'fa', 'bp'):                    # bp 最后画，压在最上面
    ser = cd.add_series(style[m][0])
    for ep in range(1, len(acc[m])):
        ser.add_data_point(ep, round(float(acc[m][ep]), 2))
ch = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, Inches(0.45), Inches(1.85), Inches(6.3), Inches(3.0), cd).chart
ch.has_legend = False; ch.font.name, ch.font.size, ch.font.color.rgb = FONT, Pt(12), GRAY
for ser, m in zip(ch.plots[0].series, ('shallow', 'fa', 'bp')):
    _, col, dash = style[m]
    ser.smooth = False; ser.format.line.color.rgb = col; ser.format.line.width = Pt(2)
    if dash:
        ser.format.line.dash_style = MSO_LINE.DASH
    ser.marker.style = XL_MARKER_STYLE.CIRCLE; ser.marker.size = 5
    ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = col; ser.marker.format.line.color.rgb = col
xa, ya = ch.category_axis, ch.value_axis
xa.minimum_scale, xa.maximum_scale, xa.major_unit = 0, 20, 5
ya.minimum_scale, ya.maximum_scale, ya.major_unit = 80, 100, 5
ya.has_major_gridlines = True; ya.major_gridlines.format.line.color.rgb = RGBColor(0xE6, 0xE6, 0xE6)
xa.has_major_gridlines = False
for ax, lab in [(xa, 'epoch'), (ya, 'accuracy (%)')]:
    ax.format.line.color.rgb = GRAY
    ax.has_title = True; ax.axis_title.text_frame.text = lab
    r = ax.axis_title.text_frame.paragraphs[0].runs[0]
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(12), False, GRAY
# 右边：图例 + 最终正确率（代替图表自带的图例）
for k, m in enumerate(('bp', 'fa', 'shallow')):
    name, col, dash = style[m]
    y = 1.75 + k * 0.6
    l = line(s, 7.3, y, 7.8, y, color=col, width=2.5)
    if dash:
        l.line.dash_style = MSO_LINE.DASH
    text(s, 8.0, y - 0.22, 3.2, 0.45, name, size=20, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 11.2, y - 0.22, 1.5, 0.45, f'{acc[m][-1]:.2f}%', size=20, bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
text(s, 7.3, 3.5, 5.4, 0.4, 'Last layer only: hidden layer never learns, only W1 is trained', size=14, color=GRAY)
text(s, 7.3, 4.1, 5.4, 0.8, 'With a random B, the hidden layer still learns', size=20, bold=True)
# 下面：同一组测试图，BP 和 FA 两行都显示（只去掉标题行，压掉图和文字之间的空白）
top, rows2 = g.crop((0, 140, g.width, 560)), g.crop((0, 665, g.width, 830))
crop = Image.new('RGB', (g.width, top.height + rows2.height), 'white')
crop.paste(top, (0, 0)); crop.paste(rows2, (0, top.height))
crop.save(BASE / 'Figures' / 'slide_guesses_bp_fa.png')
gw = 9.6
s.shapes.add_picture(str(BASE / 'Figures' / 'slide_guesses_bp_fa.png'), Inches((13.333 - gw) / 2), Inches(5.15), width=Inches(gw))
s.notes_slide.notes_text_frame.text = (
    'Here is the result. With a fixed random B, the network still learns: 98.01% correct on the test set, '
    'almost the same as backprop, 98.33%.\n'
    'But maybe only the last layer is doing the work? To check, we trained a third network where the hidden layer '
    'never learns and only W1 is trained. It reaches only 89.59%. So with a random B, the hidden layer really does learn.\n'
    'On the same test images, feedback alignment gets most of them right too. It misses this 5, and the same two hard ones '
    'that backprop also misses.\n'
    'We ran each method once, so the small gap between backprop and feedback alignment may change with a different random start.')

# ---------- 第 11 页：为什么随机的 B 也行（夹角） ----------
ang = np.load(BASE / 'Results' / 'result_fa.npz')['ang']      # 每行：(第几步, 平均夹角)
s = prs.slides.add_slide(BLANK)
title(s, 'Why Does a Random B Work?')
text(s, 0.6, 1.45, 6.2, 0.4, 'Angle between FA and BP error signals', size=18, bold=True)
cd = XyChartData()
ref = cd.add_series('90 degrees')                   # 90° 参考线
ref.add_data_point(0, 90); ref.add_data_point(20, 90)
ser = cd.add_series('angle')
for step, a in ang:
    ser.add_data_point(round(step / 60000, 4), round(float(a), 1))   # 横轴换成 epoch
ch = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(0.45), Inches(1.85), Inches(6.3), Inches(3.9), cd).chart
ch.has_legend = False; ch.font.name, ch.font.size, ch.font.color.rgb = FONT, Pt(12), GRAY
r_ser, a_ser = ch.plots[0].series
r_ser.smooth = a_ser.smooth = False
r_ser.format.line.color.rgb = GRAY; r_ser.format.line.width = Pt(1); r_ser.format.line.dash_style = MSO_LINE.ROUND_DOT
a_ser.format.line.color.rgb = BLACK; a_ser.format.line.width = Pt(2)
xa, ya = ch.category_axis, ch.value_axis
xa.minimum_scale, xa.maximum_scale, xa.major_unit = 0, 20, 5
ya.minimum_scale, ya.maximum_scale, ya.major_unit = 0, 100, 20
ya.has_major_gridlines = True; ya.major_gridlines.format.line.color.rgb = RGBColor(0xE6, 0xE6, 0xE6)
xa.has_major_gridlines = False
for ax, lab in [(xa, 'epoch'), (ya, 'angle (degrees)')]:
    ax.format.line.color.rgb = GRAY
    ax.has_title = True; ax.axis_title.text_frame.text = lab
    r = ax.axis_title.text_frame.paragraphs[0].runs[0]
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(12), False, GRAY
text(s, 3.6, 2.05, 3.0, 0.3, '90° = unrelated', size=12, color=GRAY, align=PP_ALIGN.RIGHT)
# 右上：两个箭头和夹角
OX, OY, L = 7.8, 3.55, 2.0                          # 起点和箭头长度
import math as m_
for deg, who, lab in [(8, 'BP sends back:', 'W_1^T e'), (55, 'FA sends back:', 'B e')]:
    ex, ey = OX + L * m_.cos(m_.radians(deg)), OY - L * m_.sin(m_.radians(deg))
    line(s, OX, OY, ex, ey, width=2, arrow=True)
    text(s, ex + 0.15, ey - 0.2, 1.75, 0.4, who, size=16, anchor=MSO_ANCHOR.MIDDLE)
    math(s, ex + 1.9, ey - 0.22, 1.0, 0.45, lab, size=18, align=PP_ALIGN.LEFT)
math(s, OX + 0.5, OY - 0.6, 0.4, 0.4, 'θ', size=20)
text(s, OX, OY + 0.1, 4.5, 0.3, 'e = error at the output', size=12, color=GRAY)
# 右下：怎么读这张图
t = text(s, 7.3, 4.2, 5.4, 2.3,
         ['Start: about 90°, unrelated',
          'After 1,000 images: about 50°',
          'Under 90°: FA still pushes the weights roughly the right way',
          'W1 changes to line up with B. That is the "alignment"'], size=18, space=8)
t.text_frame.paragraphs[3].runs[0].font.bold = True
s.notes_slide.notes_text_frame.text = (
    'Why does a random B work? For each image, we compare two arrows: what backprop would send back to the hidden '
    'layer, W1 transpose times the output error, and what feedback alignment actually sends back, B times the output error. '
    'We measure the angle between them and average over 500 images.\n'
    'At the start, the angle is about 90 degrees: the two arrows are unrelated. But after only about 1,000 images it drops '
    'to about 50 degrees, and it stays below 90 for the whole training.\n'
    'Below 90 degrees means the feedback alignment update still points roughly the same way as the backprop update. '
    'It is like walking downhill: you do not need the steepest direction, any direction within 90 degrees of it still takes you down.\n'
    'Why does the angle drop? Because W1 is trained using signals that went through B, W1 slowly changes to line up with B. '
    'That is where the name feedback alignment comes from.')

# ---------- 第 12 页：结论 ----------
s = prs.slides.add_slide(BLANK)
title(s, 'Skill Demo: Conclusions')
take = [(f'{acc["bp"][-1]:.2f}%', 'Backprop',
         'Using only what we learned in class, we built our own network that reads handwritten digits'),
        (f'{acc["fa"][-1]:.2f}%', 'Feedback alignment',
         'Sending the error back through a fixed random B instead of W1ᵀ still trains the hidden layer'),
        ('< 90°', 'Why it works',
         'W1 lines up with B, so the error signal points roughly the right way')]
for k, (num, head, body) in enumerate(take):
    y = 1.75 + k * 1.45
    text(s, 0.6, y, 2.8, 0.8, num, size=40, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.6, y + 0.02, 9.1, 0.4, head, size=20, bold=True)
    text(s, 3.6, y + 0.45, 9.1, 0.8, body, size=20)
t = text(s, 0.6, 6.4, 12.1, 0.4, 'Code: ', size=16, color=GRAY)
r = t.text_frame.paragraphs[0].add_run(); r.text = 'github.com/vixtor771/handwritten-digit-recognition'
r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(16), GRAY
r.hyperlink.address = 'https://github.com/vixtor771/handwritten-digit-recognition'
s.notes_slide.notes_text_frame.text = (
    'To sum up. First, using only the chain rule we learned in class, we built our own network from scratch '
    'and it reads handwritten digits with 98.33% accuracy.\n'
    'Second, backprop needs the feedback weights to match the forward weights exactly, which is hard for the brain. '
    'Feedback alignment replaces them with a fixed random matrix B, and the network still learns: 98.01%.\n'
    'Third, it works because W1 lines up with B, so the error signal stays within 90 degrees of the backprop one.\n'
    'Our code is on GitHub. Thank you, we are happy to take questions.')

prs.save(OUT)
print('saved', OUT)
